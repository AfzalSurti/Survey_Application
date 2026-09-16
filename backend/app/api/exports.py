import json
import re
from datetime import datetime
from io import BytesIO
from uuid import UUID

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.api.records import _record_filters
from app.core.deps import get_current_user
from app.database import get_db
from app.models import SurveyRecord, SurveyStatus, User
from app.schemas.survey import SurveyRecordOut
from app.services.record_enrichment import enrich_records
from app.services.work_report import collect_photo_blobs

router = APIRouter(prefix="/api/exports", tags=["exports"])

# Already surfaced as their own columns — don't repeat them as question columns.
_SKIP_RESPONSE_KEYS = {"gps", "capturedAt", "structure_category", "photos"}
_PHOTO_COLUMNS = 4  # matches the field app's minimum photo count


def _label(key: str) -> str:
    return (key or "").replace("_", " ").strip().title()


def _cell_value(value: object) -> object:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return json.dumps(value, ensure_ascii=False)


def _safe_sheet_name(name: str, used: set[str]) -> str:
    cleaned = re.sub(r"[\[\]\*\?/\\:]", "", name).strip()[:28] or "Structure"
    candidate = cleaned
    i = 2
    while candidate in used:
        candidate = f"{cleaned}-{i}"
        i += 1
    used.add(candidate)
    return candidate


def _write_overview(workbook: Workbook, records: list[SurveyRecordOut]) -> None:
    sheet = workbook.active
    sheet.title = "Overview"
    sheet.append(
        [
            "Project Name",
            "Project No.",
            "Survey Type",
            "Key Person / Highway Engineer",
            "Head Survey Person",
            "Assign Date",
            "Complete Date",
            "Chainage",
            "Category",
            "Status",
            "Captured At",
            "Latitude",
            "Longitude",
        ]
    )
    for record in records:
        sheet.append(
            [
                record.project_name,
                record.project_number,
                record.survey_type,
                record.key_engineer_name,
                record.head_surveyor_name,
                record.assign_date.isoformat() if record.assign_date else None,
                record.complete_date.isoformat() if record.complete_date else None,
                record.chainage,
                record.structure_category,
                record.status.value,
                record.captured_at.isoformat() if record.captured_at else None,
                float(record.latitude) if record.latitude is not None else None,
                float(record.longitude) if record.longitude is not None else None,
            ]
        )
    for col in range(1, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 20


def _write_category_sheet(
    workbook: Workbook,
    category: str,
    records: list[SurveyRecordOut],
    photo_blobs: dict[UUID, list[bytes]],
    used_names: set[str],
) -> None:
    # Every question that appears anywhere in this category's answers becomes
    # its own column — different categories ask different questions, so the
    # column set is built from the data rather than hard-coded.
    question_keys: list[str] = []
    seen: set[str] = set()
    for record in records:
        for key in (record.responses_json or {}).keys():
            if key in _SKIP_RESPONSE_KEYS or key in seen:
                continue
            seen.add(key)
            question_keys.append(key)

    headers = (
        ["Project", "Chainage", "Surveyor", "Status", "Captured At", "Latitude", "Longitude"]
        + [_label(k) for k in question_keys]
        + [f"Photo {i + 1}" for i in range(_PHOTO_COLUMNS)]
    )
    sheet = workbook.create_sheet(title=_safe_sheet_name(_label(category) or "Structure", used_names))
    sheet.append(headers)
    photo_col_start = len(headers) - _PHOTO_COLUMNS + 1

    for row_idx, record in enumerate(records, start=2):
        responses = record.responses_json or {}
        row = [
            record.project_name,
            record.chainage,
            record.head_surveyor_name,
            record.status.value,
            record.captured_at.isoformat() if record.captured_at else None,
            float(record.latitude) if record.latitude is not None else None,
            float(record.longitude) if record.longitude is not None else None,
        ] + [_cell_value(responses.get(k)) for k in question_keys]
        sheet.append(row + [None] * _PHOTO_COLUMNS)

        blobs = photo_blobs.get(record.id, [])
        if blobs:
            sheet.row_dimensions[row_idx].height = 60
        for i, blob in enumerate(blobs[:_PHOTO_COLUMNS]):
            try:
                image = XLImage(BytesIO(blob))
                image.width, image.height = 80, 60
                image.anchor = f"{get_column_letter(photo_col_start + i)}{row_idx}"
                sheet.add_image(image)
            except Exception:
                continue

    for col in range(1, photo_col_start):
        sheet.column_dimensions[get_column_letter(col)].width = 18
    for col in range(photo_col_start, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 12


@router.get("/excel")
async def export_excel(
    project_id: UUID | None = None,
    chainage: str | None = None,
    structure_category: str | None = None,
    status: SurveyStatus | None = None,
    surveyor_id: UUID | None = None,
    date_from: datetime | None = None,
    date_to: datetime | None = None,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
) -> StreamingResponse:
    filters = _record_filters(
        project_id, chainage, structure_category, status, surveyor_id, date_from, date_to, user
    )
    result = await db.execute(
        select(SurveyRecord)
        .where(*filters)
        .options(selectinload(SurveyRecord.photos))
        .order_by(SurveyRecord.captured_at.desc())
    )
    raw_records = list(result.scalars().all())
    enriched = await enrich_records(db, raw_records)
    photo_blobs = await collect_photo_blobs(raw_records)

    workbook = Workbook()
    _write_overview(workbook, enriched)

    by_category: dict[str, list[SurveyRecordOut]] = {}
    for record in enriched:
        by_category.setdefault(record.structure_category, []).append(record)

    used_names = {"Overview"}
    for category, records in by_category.items():
        _write_category_sheet(workbook, category, records, photo_blobs, used_names)

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="gdrpl-survey-records.xlsx"'},
    )
